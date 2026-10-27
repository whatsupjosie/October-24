"""Real format and auth tests: no Office application or cloud service required."""

import base64
import io
import zipfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from modules.foresight_file_formats import create_file_format_router, convert
from modules.auth import create_access_token


def docx_bytes(text):
    doc = Document()
    doc.add_paragraph(text)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table content"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_docx_read_original_unchanged_and_export_roundtrip(monkeypatch):
    monkeypatch.setenv("PUBCAST_ENFORCE_AUTH", "1")
    app = FastAPI()
    app.include_router(create_file_format_router())
    client = TestClient(app)
    raw = docx_bytes("Hello\tCurtsey\nSecond line")
    original = bytes(raw)
    body = {"name": "sample.docx", "data": base64.b64encode(raw).decode()}
    assert client.post("/api/foresight/files/convert", json=body).status_code == 401
    headers = {
        "Authorization": "Bearer "
        + create_access_token({"sub": "owner", "role": "owner"})
    }
    result = client.post("/api/foresight/files/convert", headers=headers, json=body)
    assert result.status_code == 200, result.text
    assert result.json()["text"] == "Hello\tCurtsey\nSecond line\nTable content"
    assert "not preserved" in result.json()["warning"]
    exported = client.post(
        "/api/foresight/files/export",
        headers=headers,
        json={"text": result.json()["text"] + "\nEdited", "format": "docx"},
    )
    assert exported.status_code == 200
    assert convert("output.docx", exported.content)["text"].endswith("\nEdited")
    assert raw == original


def test_xlsx_formulas_are_text_not_false_calculated_values():
    book = Workbook()
    book.active.append(["Name", "Formula"])
    book.active.append(["Curtsey", "=1+2"])
    data = io.BytesIO()
    book.save(data)
    result = convert("sheet.xlsx", data.getvalue())
    assert "Curtsey\t=1+2" in result["text"]
    assert "not recalculated" in result["warning"]


def test_invalid_and_entity_archives_rejected():
    with pytest.raises(ValueError, match="Invalid document archive"):
        convert("bad.docx", b"not a zip")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", '<!DOCTYPE foo [<!ENTITY x "oops">]><foo/>')
    with pytest.raises(ValueError, match="entities"):
        convert("entity.docx", out.getvalue())


def test_odt_and_pptx_extraction():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr(
            "content.xml",
            '<root xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><text:p>One <text:span>two</text:span></text:p></root>',
        )
    assert convert("test.odt", out.getvalue())["text"] == "One two"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for number in [10, 2, 1]:
            z.writestr(
                f"ppt/slides/slide{number}.xml",
                f'<root xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>{number}</a:t></root>',
            )
    assert (
        convert("test.pptx", out.getvalue())["text"]
        == "--- Slide 1 ---\n1\n--- Slide 2 ---\n2\n--- Slide 3 ---\n10"
    )


def test_pdf_text_extraction_without_office_or_ocr():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 10 Td (Curtsey PDF) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    out = io.BytesIO()
    writer.write(out)
    result = convert("text.pdf", out.getvalue())
    assert "Curtsey PDF" in result["text"]
    assert "OCR" in result["warning"]


def test_native_docx_export_does_not_require_external_writer(monkeypatch):
    import sys

    monkeypatch.setenv("PUBCAST_ENFORCE_AUTH", "0")
    monkeypatch.setitem(sys.modules, "docx", None)
    app = FastAPI()
    app.include_router(create_file_format_router())
    client = TestClient(app)
    result = client.post(
        "/api/foresight/files/export",
        json={"text": "  Spaces\t<&>\nUnicode: café 🌳\n", "format": "docx"},
    )
    assert result.status_code == 200
    assert (
        convert("native.docx", result.content)["text"]
        == "  Spaces\t<&>\nUnicode: café 🌳\n"
    )
    # Independent reader verifies package structure, not merely our own parser.
    assert Document(io.BytesIO(result.content)).paragraphs[0].text == "  Spaces\t<&>"
    assert (
        client.post(
            "/api/foresight/files/export",
            json={"text": "bad\u0000text", "format": "docx"},
        ).status_code
        == 422
    )


def test_duplicate_document_members_are_rejected():
    import warnings

    out = io.BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with zipfile.ZipFile(out, "w") as z:
            z.writestr("word/document.xml", "<document/>")
            z.writestr("word/document.xml", "<document/>")
    with pytest.raises(ValueError, match="duplicate member names"):
        convert("ambiguous.docx", out.getvalue())
