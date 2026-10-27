"""Offline, byte-level acceptance for versioned rich .2i persistence.

These checks qualify portable storage and authenticated local API only. They do
not claim the rich editor has been mounted inside production 2i yet.
"""
import base64
import hashlib
import io
import json
import uuid
import warnings
import zipfile
from dataclasses import replace

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from modules.auth import create_access_token
from modules.foresight_file_formats import create_file_format_router
from modules.foresight_workspace import create_workspace_router
from modules.twoi_rich_document import (
    FORMAT, MANIFEST, ORIGINAL, WORKING, RichDocument, RichDocumentError,
    identify, pack, unpack,
)


def document_bytes():
    doc = Document()
    doc.add_heading("Preserved", 1)
    run = doc.add_paragraph().add_run("Formatting stays")
    run.bold = True
    run.italic = True
    run.font.name = "Liberation Serif"
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Original table"
    doc.sections[0].header.paragraphs[0].text = "HEADER"
    doc.sections[0].footer.paragraphs[0].text = "FOOTER"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def rich():
    source = document_bytes()
    new = Document(io.BytesIO(source))
    new.add_paragraph("Working revision two")
    stream = io.BytesIO()
    new.save(stream)
    return RichDocument(
        document_id=str(uuid.UUID("1a0cad75-b330-4b74-837a-ea0cfa9d9880")),
        original_name="Manuscript ✨.docx",
        original_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        original=source,
        working_docx=stream.getvalue(),
        metadata={"title": "Lyrical Manuscript", "future-field": {"keep": [1, "two"]}},
    )


def rezip(raw, mutations):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w") as dest:
        for name in src.namelist():
            if name not in mutations:
                dest.writestr(name, src.read(name))
            elif mutations[name] is not None:
                dest.writestr(name, mutations[name])
        for name, value in mutations.items():
            if name not in src.namelist() and value is not None:
                dest.writestr(name, value)
    return out.getvalue()


def test_deterministic_roundtrip_retains_exact_source_and_working_bytes(rich):
    before = bytes(rich.original)
    packed = pack(rich)
    assert packed == pack(rich)
    assert identify(packed) == "rich-v5"
    result = unpack(packed)
    assert result == rich
    assert result.original == before
    assert hashlib.sha256(result.original).digest() == hashlib.sha256(before).digest()
    doc = Document(io.BytesIO(result.working_docx))
    assert doc.paragraphs[-1].text == "Working revision two"
    assert doc.paragraphs[1].runs[0].bold and doc.paragraphs[1].runs[0].italic
    assert doc.tables[0].cell(0, 0).text == "Original table"
    assert doc.sections[0].header.paragraphs[0].text == "HEADER"
    assert doc.sections[0].footer.paragraphs[0].text == "FOOTER"
    assert result.metadata["future-field"] == {"keep": [1, "two"]}


def test_legacy_version_identified_not_migrated_or_rewritten():
    old = b'{"format":"2i-document","version":4,"state":{},"unknown":true}'
    assert identify(old) == "manuscript-v4"
    with pytest.raises(RichDocumentError, match="archive"):
        unpack(old)
    with pytest.raises(RichDocumentError, match="Unsupported"):
        identify(b'{"format":"2i-document","version":9,"state":{}}')


def test_checksum_mismatch_is_fatal(rich):
    old = pack(rich)
    tampered = rezip(old, {ORIGINAL: rich.original + b"tampered"})
    with pytest.raises(RichDocumentError, match="integrity"):
        unpack(tampered)


def test_missing_extra_and_duplicate_members_rejected(rich):
    raw = pack(rich)
    for wrong in [rezip(raw, {WORKING: None}), rezip(raw, {"../../bad.txt": b"x"})]:
        with pytest.raises(RichDocumentError, match="members"):
            unpack(wrong)
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as z, zipfile.ZipFile(buffer, "w") as out:
        for p in z.namelist():
            out.writestr(p, z.read(p))
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Duplicate name:")
            out.writestr(MANIFEST, z.read(MANIFEST))
    with pytest.raises(RichDocumentError, match="members"):
        unpack(buffer.getvalue())


def test_future_version_rejected_without_guessing(rich):
    raw = pack(rich)
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        manifest = json.loads(z.read(MANIFEST))
    manifest["version"] += 1
    wrong = rezip(raw, {MANIFEST: json.dumps(manifest).encode()})
    with pytest.raises(RichDocumentError, match="Unsupported"):
        unpack(wrong)


@pytest.mark.parametrize("field,value,error", [
    ("document_id", "random", "UUID"),
    ("original_name", "../bad.docx", "path"),
    ("original_name", "line\nfeed.docx", "control"),
    ("original_mime", "bad\n", "MIME"),
    ("metadata", {"invalid": float("nan")}, "metadata"),
    ("original", b"X" * (24 * 1024 * 1024 + 1), "Original"),
    ("working_docx", b"not docx", "DOCX"),
], ids=["invalid-document-id", "unsafe-name", "control-in-name", "invalid-mime",
        "nonfinite-metadata", "oversized-original", "invalid-working-docx"])
def test_invalid_material_never_packages(rich, field, value, error):
    with pytest.raises(RichDocumentError, match=error):
        pack(replace(rich, **{field: value}))


def test_manifest_rejects_duplicate_json_fields(rich):
    raw = pack(rich)
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        manifest = z.read(MANIFEST)
    new_manifest = manifest[:-1] + b',"version":5}'
    with pytest.raises(RichDocumentError, match="manifest"):
        unpack(rezip(raw, {MANIFEST: new_manifest}))


def test_api_enforces_auth_then_roundtrips_with_workspace_and_revision_guards(monkeypatch, tmp_path, rich):
    monkeypatch.setenv("PUBCAST_ENFORCE_AUTH", "1")
    app = FastAPI()
    app.include_router(create_file_format_router())
    app.include_router(create_workspace_router(tmp_path))
    client = TestClient(app)
    token = create_access_token({"sub": "2i-owner", "role": "owner"})
    auth = {"Authorization": "Bearer " + token}
    body = {
        "document_id": rich.document_id,
        "original_name": rich.original_name,
        "original_mime": rich.original_mime,
        "original_data_b64": base64.b64encode(rich.original).decode(),
        "working_docx_b64": base64.b64encode(rich.working_docx).decode(),
        "metadata": rich.metadata,
    }
    path = "/api/foresight/files/native-rich/create"
    assert client.post(path, json=body).status_code == 401
    result = client.post(path, json=body, headers=auth)
    assert result.status_code == 200, result.text
    bundle = base64.b64decode(result.json()["data_b64"])
    assert hashlib.sha256(bundle).hexdigest() == result.json()["sha256"]
    destination = "/api/foresight/workspace/file"
    save = client.post(destination, headers=auth, json={
        "path": "Manuscripts/Manuscript V2.2i", "data": result.json()["data_b64"],
        "expected_revision": None,
    })
    assert save.status_code == 200, save.text
    stale = client.post(destination, headers=auth, json={
        "path": "Manuscripts/Manuscript V2.2i", "data": result.json()["data_b64"],
        "expected_revision": "stale",
    })
    assert stale.status_code in (409, 412)
    read = client.get(destination, headers=auth, params={"path": "Manuscripts/Manuscript V2.2i"})
    assert read.status_code == 200, read.text
    restored_bundle = base64.b64decode(read.json()["data"])
    assert restored_bundle == bundle
    restored = client.post("/api/foresight/files/native-rich/read", headers=auth,
                           json={"data_b64": base64.b64encode(restored_bundle).decode()})
    assert restored.status_code == 200, restored.text
    content = restored.json()
    assert content["original_sha256"] == hashlib.sha256(rich.original).hexdigest()
    assert base64.b64decode(content["original_data_b64"]) == rich.original
    assert base64.b64decode(content["working_docx_b64"]) == rich.working_docx
    assert content["metadata"] == rich.metadata
    bad = client.post("/api/foresight/files/native-rich/read", headers=auth,
                      json={"data_b64": "not base64**"})
    assert bad.status_code == 422


def test_invalid_zip_bytes_do_not_escape_as_server_error():
    with pytest.raises(RichDocumentError, match="archive"):
        unpack(b"PK\x03\x04clearly not a real zip")


def test_symlinks_in_portable_container_rejected(rich):
    raw = pack(rich)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(output, "w") as dst:
        for name in src.namelist():
            info = zipfile.ZipInfo(name)
            if name == ORIGINAL:
                info.create_system = 3
                info.external_attr = 0o120777 << 16
            dst.writestr(info, src.read(name))
    with pytest.raises(RichDocumentError, match="Symbolic links"):
        unpack(output.getvalue())


def test_invalid_embedded_docx_paths_not_accepted(rich):
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(rich.working_docx)) as src, zipfile.ZipFile(stream, "w") as dest:
        for name in src.namelist():
            dest.writestr(name, src.read(name))
        dest.writestr("../outside.bin", b"no")
    with pytest.raises(RichDocumentError, match="Unsafe DOCX member path"):
        pack(replace(rich, working_docx=stream.getvalue()))
