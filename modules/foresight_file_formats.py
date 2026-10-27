"""Local, non-destructive document adapters for tray working copies.

No file paths or remote URLs are accepted. Original bytes stay with the client.
Conversion explicitly returns editable text, not a fidelity-preserving Office editor.
"""

from __future__ import annotations

import base64
import binascii
import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from .route_security import current_identity
from .twoi_rich_document import (
    RichDocument, RichDocumentError, MAX_CONTAINER, MAX_PART,
    pack as pack_rich, unpack as unpack_rich,
)

MAX_BYTES = 24 * 1024 * 1024
MAX_EXPANDED = 64 * 1024 * 1024
MAX_TEXT = 5 * 1024 * 1024


class ConvertRequest(BaseModel):
    name: str = Field(max_length=255)
    data: str = Field(max_length=MAX_BYTES * 4 // 3 + 8)


class ExportRequest(BaseModel):
    text: str = Field(max_length=MAX_TEXT)
    format: str


class RichPackRequest(BaseModel):
    document_id: str = Field(max_length=36)
    original_name: str = Field(max_length=255)
    original_mime: str = Field(max_length=255)
    original_data_b64: str = Field(max_length=MAX_PART * 4 // 3 + 8)
    working_docx_b64: str = Field(max_length=MAX_PART * 4 // 3 + 8)
    metadata: dict = Field(default_factory=dict)


class RichUnpackRequest(BaseModel):
    data_b64: str = Field(max_length=MAX_CONTAINER * 4 // 3 + 8)


def archive(raw: bytes) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(io.BytesIO(raw))
        entries = z.infolist()
        if len(entries) > 10000 or sum(e.file_size for e in entries) > MAX_EXPANDED:
            z.close()
            raise ValueError("Document archive exceeds the conversion limit")
        if any(e.flag_bits & 1 for e in entries):
            z.close()
            raise ValueError("Encrypted documents must be unlocked before conversion")
        # ZIP readers otherwise resolve duplicate names inconsistently; reject
        # ambiguous documents rather than interpreting attacker-selected parts.
        names = [entry.filename for entry in entries]
        if len(names) != len(set(names)):
            z.close()
            raise ValueError("Document archive contains duplicate member names")
        return z
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid document archive") from exc


def xml_member(z: zipfile.ZipFile, name: str):
    data = z.read(name)
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValueError("Document XML entities are not supported")
    return ET.fromstring(data)


def convert(name: str, raw: bytes) -> dict:
    if len(raw) > MAX_BYTES:
        raise ValueError("Document exceeds 24 MB conversion limit")
    ext = Path(name).suffix.lower()
    warning = "Editable text copy: layout, styling, images, comments and other document features are not preserved. The original is unchanged."
    if ext in {".docx", ".odt", ".pptx", ".xlsx"}:
        with archive(raw) as z:
            if ext == ".docx":
                warning += " Main document text and table cells are extracted; headers, footers and footnotes are not included."
                ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                root = xml_member(z, "word/document.xml")
                lines = []
                for p in root.iter(ns + "p"):
                    parts = []
                    for node in p.iter():
                        if node.tag == ns + "t":
                            parts.append(node.text or "")
                        elif node.tag == ns + "tab":
                            parts.append("\t")
                        elif node.tag in {ns + "br", ns + "cr"}:
                            parts.append("\n")
                    lines.append("".join(parts))
            elif ext == ".odt":
                root = xml_member(z, "content.xml")
                ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
                lines = [
                    "".join(p.itertext())
                    for p in root.iter()
                    if p.tag in {ns + "p", ns + "h"}
                ]
            elif ext == ".pptx":
                import re

                names = sorted(
                    (
                        n
                        for n in z.namelist()
                        if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
                    ),
                    key=lambda n: int(re.search(r"(\d+)\.xml", n)[1]),
                )
                ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                lines = []
                for i, n in enumerate(names, 1):
                    lines.append(f"--- Slide {i} ---")
                    lines.extend(
                        "".join(t.itertext()) for t in xml_member(z, n).iter(ns + "t")
                    )
                warning += " Slide positioning and speaker notes are not included."
            else:
                from openpyxl import load_workbook

                book = load_workbook(
                    io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
                )
                lines = []
                try:
                    count = 0
                    for sheet in book:
                        lines.append("--- " + sheet.title + " ---")
                        for row in sheet.iter_rows(values_only=True):
                            count += len(row)
                            if count > 200000:
                                raise ValueError(
                                    "Spreadsheet exceeds 200,000-cell conversion limit"
                                )
                            lines.append(
                                "\t".join("" if v is None else str(v) for v in row)
                            )
                finally:
                    book.close()
                warning += " Cells are tab-separated; formulas are shown as text and are not recalculated."
            text = "\n".join(lines)
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs must be unlocked before conversion")
        if len(reader.pages) > 500:
            raise ValueError("PDF exceeds 500-page conversion limit")
        text = "\n\n".join(p.extract_text() or "" for p in reader.pages)
        warning = "Extracted PDF text; layout is not retained. Scanned pages require OCR, which is not provided here. Original unchanged."
    else:
        raise ValueError(
            "No local converter for this format. Keep the original or use an installed format-specific application."
        )
    if len(text) > MAX_TEXT:
        raise ValueError("Extracted text exceeds 5 MB editor limit")
    return {"text": text, "warning": warning, "format": ext.lstrip(".")}


def create_file_format_router():
    router = APIRouter(
        prefix="/api/foresight/files",
        tags=["Foresight files"],
        dependencies=[Depends(current_identity)],
    )

    @router.post("/native-rich/create")
    def create_rich_working_copy(body: RichPackRequest):
        """Offline-only byte-preserving v5 container; never modifies originals."""
        from hashlib import sha256

        try:
            document = RichDocument(
                document_id=body.document_id,
                original_name=body.original_name,
                original_mime=body.original_mime,
                original=base64.b64decode(body.original_data_b64, validate=True),
                working_docx=base64.b64decode(body.working_docx_b64, validate=True),
                metadata=body.metadata,
            )
            raw = pack_rich(document)
        except (RichDocumentError, binascii.Error) as exc:
            raise HTTPException(422, str(exc)) from exc
        return {
            "kind": "rich-v5",
            "data_b64": base64.b64encode(raw).decode("ascii"),
            "sha256": sha256(raw).hexdigest(),
            "size": len(raw),
        }

    @router.post("/native-rich/read")
    def read_rich_working_copy(body: RichUnpackRequest):
        """Verify contents and return source and working bytes separately."""
        from hashlib import sha256

        try:
            raw = base64.b64decode(body.data_b64, validate=True)
            document = unpack_rich(raw)
        except (RichDocumentError, binascii.Error) as exc:
            raise HTTPException(422, str(exc)) from exc
        return {
            "kind": "rich-v5",
            "document_id": document.document_id,
            "original_name": document.original_name,
            "original_mime": document.original_mime,
            "original_data_b64": base64.b64encode(document.original).decode("ascii"),
            "original_sha256": sha256(document.original).hexdigest(),
            "working_docx_b64": base64.b64encode(document.working_docx).decode("ascii"),
            "working_sha256": sha256(document.working_docx).hexdigest(),
            "metadata": document.metadata,
        }

    @router.post("/convert")
    def convert_file(body: ConvertRequest):
        try:
            raw = base64.b64decode(body.data, validate=True)
            return convert(body.name, raw)
        except ImportError as exc:
            raise HTTPException(
                503, "Install the declared document dependencies to use this converter"
            ) from exc
        except (
            ValueError,
            KeyError,
            ET.ParseError,
            binascii.Error,
            zipfile.BadZipFile,
            RuntimeError,
        ) as exc:
            raise HTTPException(422, str(exc)) from exc
        except Exception as exc:
            import logging

            logging.getLogger(__name__).warning(
                "Document parser rejected %s: %s", body.name, type(exc).__name__
            )
            raise HTTPException(
                422,
                "This document could not be converted. It may be damaged, encrypted or unsupported; original unchanged.",
            ) from exc

    @router.post("/export")
    def export_file(body: ExportRequest):
        if body.format != "docx":
            raise HTTPException(422, "Supported document export: docx")
        if any(
            (ord(c) < 32 and c not in "\t\n\r")
            or 0xD800 <= ord(c) <= 0xDFFF
            or ord(c) in {0xFFFE, 0xFFFF}
            for c in body.text
        ):
            raise HTTPException(
                422, "Text contains characters not supported by XML documents"
            )
        # DOCX is an open ZIP/XML format. Plain-text copies need no Office application
        # or third-party document writer at runtime.
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        ET.register_namespace("w", ns)
        document = ET.Element("{" + ns + "}document")
        content = ET.SubElement(document, "{" + ns + "}body")
        for line in body.text.split("\n"):
            paragraph = ET.SubElement(content, "{" + ns + "}p")
            for index, part in enumerate(line.split("\t")):
                run = ET.SubElement(paragraph, "{" + ns + "}r")
                if index:
                    ET.SubElement(run, "{" + ns + "}tab")
                text = ET.SubElement(
                    run,
                    "{" + ns + "}t",
                    {"{http://www.w3.org/XML/1998/namespace}space": "preserve"},
                )
                text.text = part
        ET.SubElement(content, "{" + ns + "}sectPr")
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(
                "[Content_Types].xml",
                '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
            )
            z.writestr(
                "_rels/.rels",
                '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
            )
            z.writestr(
                "word/document.xml",
                ET.tostring(document, encoding="utf-8", xml_declaration=True),
            )
        from fastapi.responses import Response

        return Response(
            out.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    return router
