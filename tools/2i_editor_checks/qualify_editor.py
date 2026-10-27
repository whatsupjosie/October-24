from pathlib import Path
from playwright.sync_api import sync_playwright
from urllib.parse import urlparse
from docx import Document
import json, tempfile, subprocess
from docx.shared import Inches, Pt
root=Path(__file__).resolve().parents[2]
work=Path(tempfile.mkdtemp(prefix='twoi-editor-qualification-'))
base=root/'2i/editor-engine/node_modules'
subprocess.run([str(base/'.bin/esbuild'),str(root/'2i/editor-engine/qualification-entry.js'),'--bundle','--platform=browser','--format=esm','--outfile='+str(work/'editor-qualification.js')],check=True)
d=Document();d.sections[0].left_margin=Inches(1.5)
p=d.add_paragraph();r=p.add_run('INT. STUDIO - NIGHT');r.bold=True;r.font.name='Courier New';r.font.size=Pt(12)
p=d.add_paragraph('Keep this formatting.');p.paragraph_format.left_indent=Inches(.5)
t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Scene';t.cell(0,1).text='One'
d.sections[0].header.paragraphs[0].text='MANUSCRIPT V2'
d.save(work/'editor-test-original.docx')
requests=[];errors=[]
with sync_playwright() as p:
 browser=p.chromium.launch(args=['--single-process','--no-zygote','--use-angle=swiftshader'])
 page=browser.new_page(viewport={'width':1200,'height':900})
 page.on('pageerror',lambda e:errors.append(str(e)))
 def route(r):
  url=urlparse(r.request.url)
  if url.netloc!='localhost':requests.append(r.request.url);r.abort();return
  path=url.path
  files={'/harfbuzz.wasm':base/'harfbuzzjs/dist/harfbuzz.wasm','/engine.js':work/'editor-qualification.js','/editor.css':base/'@docx-editor.dev/core/dist/editor.css','/local-font.ttf':base/'@docx-editor.dev/fonts/assets/LiberationMono-Regular.ttf'}
  if path=='/':r.fulfill(content_type='text/html',body='<html><head><link rel="stylesheet" href="/editor.css"></head><body><div class="docx-editor" id="editor"></div><script type="module" src="/engine.js"></script></body></html>')
  elif path in files:r.fulfill(path=str(files[path]))
  else:requests.append(path);r.abort()
 page.route('**/*',route)
 page.goto('http://localhost/')
 page.wait_for_function('typeof window.qualifyEngine === "function"',timeout=20000)
 page.evaluate('bytes=>qualifyEngine(bytes)',list((work/'editor-test-original.docx').read_bytes()))
 page.wait_for_function('window.editor && !editor.snapshot().isLoading && !editor.snapshot().isOpening',timeout=20000)
 state=page.evaluate('({editable:editor.snapshot().editable,error:editor.snapshot().parseError,matches:editor.findMatches("Keep this formatting.").length})')
 assert state['editable'] and state['error'] is None and state['matches']==1,state
 result=page.evaluate("()=>{const [match]=editor.findMatches('Keep this formatting.');return editor.exec({type:'replaceMatch',match,text:'Preserve this formatting.'})}")
 assert result['ok'],result
 saved=page.evaluate('async()=>Array.from(new Uint8Array(await editor.save()))')
 out=work/'editor-test-saved.docx';out.write_bytes(bytes(saved))
 page.evaluate('bytes=>editor.load(new Uint8Array(bytes))',saved)
 page.wait_for_function('!editor.snapshot().isLoading && !editor.snapshot().isOpening')
 assert page.evaluate('editor.findMatches("Preserve this formatting.").length')==1
 assert not requests,requests
 assert not errors,errors
 browser.close()
d=Document(out)
assert d.paragraphs[0].text=='INT. STUDIO - NIGHT'
r=d.paragraphs[0].runs[0]
assert r.bold and r.font.name=='Courier New' and r.font.size.pt==12
assert d.paragraphs[1].text=='Preserve this formatting.'
assert abs(d.paragraphs[1].paragraph_format.left_indent.inches-.5)<.001
assert d.tables[0].cell(0,1).text=='One'
assert d.sections[0].header.paragraphs[0].text=='MANUSCRIPT V2'
assert abs(d.sections[0].left_margin.inches-1.5)<.001
print(json.dumps({'offline_engine_open_edit_save_reopen':'passed','bold_font_size_indent_header_table_margin':'preserved','external_requests':requests,'browser_errors':errors}))
